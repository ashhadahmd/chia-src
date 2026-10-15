"""Render a Chia doc (editor HTML) to PDF or DOCX, a sheet to XLSX and slides to PPTX.

The editor emits a small, known subset of HTML: p, h1-h3, ul/ol/li, blockquote, hr, br,
strong/em/u/s/code, <span style="font-size; color">, `text-align` on blocks, tables
(with colspan/rowspan and column widths), images (<img data-width data-align>) and diagrams
(<figure data-type="diagram"> holding a rendered <img>). Both renderers understand exactly
that subset. Images are resolved from Chia's asset store; nothing is fetched from the web.
"""

import io
import re
from html import escape
from pathlib import Path
from typing import Any

from app.assets import Resolver, inline
from app.models import PageSize

FONT_DIR = (Path(__file__).parent / "assets" / "fonts").resolve()
HEX = re.compile(r"^#(?:[0-9a-fA-F]{3}){1,2}$")
INK_LIGHT, INK_DARK = "#2C2C29", "#ECE7DC"
HEADER_FILL = "EDEDED"  # table header rows: light grey with black text

# Text colours are stored as `var(--tc-<name>, <light hex>)` so the editor can adapt them to
# its theme. PDF renders resolve them for the chosen page colour; Word/Google get the hex.
TEXT_COLORS = {
    "stone": ("#6F6E66", "#A9A69B"),
    "sage": ("#5A6B47", "#A3B68F"),
    "eucalyptus": ("#3F6E66", "#8BB5AC"),
    "terracotta": ("#B05A47", "#DC9A8A"),
    "ochre": ("#9A7424", "#D9B562"),
    "clay": ("#8C5A3C", "#C99A7C"),
    "plum": ("#7A4F72", "#C79BC0"),
    "slate": ("#4A5A6E", "#9DB0C6"),
}
_VAR_COLOR = re.compile(r"var\(--tc-[a-z]+,\s*(#[0-9a-fA-F]{6})\)")


def resolve_colors(html: str) -> str:
    """Replace theme-aware colour variables with their light-page hex (for Word/Google)."""
    return _VAR_COLOR.sub(r"\1", html)


def _luminance(hex_color: str) -> float:
    h = hex_color.lstrip("#")
    if len(h) == 3:
        h = "".join(c * 2 for c in h)
    r, g, b = (int(h[i : i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def parse_background(value: str | None) -> str:
    if value is None or value == "":
        return "#FFFFFF"
    if not HEX.match(value):
        raise ValueError("background must be a hex colour like #FAF8F5")
    return value


def _no_resolver(_: str) -> None:
    return None


_EMPTY_DIAGRAM = re.compile(r"<figure([^>]*\bdata-type=\"diagram\"[^>]*)>\s*</figure>")
_DATA_CODE = re.compile(r'data-code="([^"]*)"')


def _diagram_fallbacks(html: str) -> str:
    """Diagrams written by an MCP client have no picture until the doc is opened in Chia;
    until then exports show their source."""

    def source(m: re.Match[str]) -> str:
        code = _DATA_CODE.search(m.group(1))
        return f"<pre>{code.group(1) if code else ''}</pre>"

    return _EMPTY_DIAGRAM.sub(source, html)


# ───────────── PDF ─────────────


class PdfUnavailable(RuntimeError):
    """WeasyPrint or its native libraries (Pango, Cairo) aren't installed, as on the slim API
    Lambda. Browsers print PDFs themselves; this export is only a fallback."""


def _weasyprint() -> tuple[Any, type]:
    """WeasyPrint is imported on first use, so the API starts without Pango/Cairo present."""
    try:
        from weasyprint import HTML
        from weasyprint.urls import URLFetcher
    except (ImportError, OSError) as exc:  # OSError: the package is there, its C libraries aren't
        raise PdfUnavailable(f"Server-side PDF export isn't available here ({exc.__class__.__name__})") from exc

    class LocalOnly(URLFetcher):
        """Documents may contain arbitrary URLs; only Chia's bundled fonts and inlined images load."""

        def fetch(self, url, headers=None):
            if url.startswith("data:image/"):
                return super().fetch(url, headers)
            if not url.startswith("file://"):
                raise ValueError(f"Blocked external resource: {url[:80]}")
            path = Path(url.removeprefix("file://")).resolve()
            if path.parent != FONT_DIR:
                raise ValueError(f"Blocked local resource: {url}")
            return super().fetch(url, headers)

    return HTML, LocalOnly


def to_pdf(title: str, html: str, page: PageSize, background: str, resolve: Resolver = _no_resolver) -> bytes:
    HTML, LocalOnly = _weasyprint()
    light_page = _luminance(background) > 0.45
    ink = INK_LIGHT if light_page else INK_DARK
    rule = "rgba(44,44,41,0.22)" if light_page else "rgba(236,231,220,0.22)"
    head = "#EDEDED" if light_page else "rgba(236,231,220,0.10)"
    head_ink = "#000000" if light_page else INK_DARK
    palette = "".join(f"--tc-{name}: {pair[0] if light_page else pair[1]};" for name, pair in TEXT_COLORS.items())
    size = "letter" if page is PageSize.letter else "A4"
    font = FONT_DIR.as_uri()
    css = f"""
    @font-face {{ font-family: Nunito; src: url({font}/nunito-latin-wght-normal.woff2); font-weight: 200 1000; }}
    @font-face {{ font-family: Nunito; font-style: italic; src: url({font}/nunito-latin-wght-italic.woff2); font-weight: 200 1000; }}
    :root {{ {palette} }}
    @page {{ size: {size}; margin: 1in; background: {background}; }}
    body {{ font-family: Nunito, sans-serif; font-size: 11pt; line-height: 1.5; color: {ink}; }}
    p, h1, h2, h3, ul, ol, blockquote {{ margin: 0; }}
    p:empty::after {{ content: "\\00a0"; }}
    h1 {{ font-size: 20pt; line-height: 1.25; margin: 0.4em 0 0.2em; }}
    h2 {{ font-size: 16pt; line-height: 1.3; margin: 0.4em 0 0.2em; }}
    h3 {{ font-size: 13pt; line-height: 1.35; margin: 0.4em 0 0.2em; }}
    ul, ol {{ padding-left: 1.6em; }}
    blockquote {{ padding-left: 1em; border-left: 3px solid currentColor; opacity: 0.85; }}
    hr {{ border: 0; border-top: 1px solid currentColor; opacity: 0.25; margin: 0.8em 0; }}
    code, pre {{ font-family: monospace; }}
    pre {{ white-space: pre-wrap; font-size: 9pt; margin: 0.5em 0; }}
    img {{ display: block; max-width: 100%; height: auto; margin: 0.5em auto; break-inside: avoid; }}
    img[data-align="left"] {{ margin-left: 0; }}
    img[data-align="right"] {{ margin-right: 0; }}
    figure {{ margin: 0.5em 0; }}
    figure img {{ width: 100%; margin: 0 auto; }}
    table {{ border-collapse: collapse; width: 100%; margin: 0.5em 0; table-layout: fixed; }}
    tr {{ break-inside: avoid; }}
    th, td {{ border: 1px solid {rule}; padding: 4pt 7pt; vertical-align: top; text-align: left; }}
    th {{ background: {head}; color: {head_ink}; font-weight: 700; }}
    thead {{ display: table-header-group; }}
    """
    body = inline(_diagram_fallbacks(html), resolve)
    document = f"<!doctype html><html><head><meta charset='utf-8'><title>{escape(title)}</title><style>{css}</style></head><body>{body}</body></html>"
    return HTML(string=document, url_fetcher=LocalOnly(allowed_protocols={"file", "data"})).write_pdf()


# ───────────── DOCX ─────────────


def to_docx(title: str, html: str, page: PageSize, resolve: Resolver = _no_resolver) -> bytes:
    # python-docx loads only when someone exports.
    from app.export_docx import to_docx as render

    return render(title, html, page, resolve)


# ───────────── Sheets and slides ─────────────


def to_xlsx(title: str, columns: list[str], rows: list[list[Any]]) -> bytes:
    """One worksheet: a grey, frozen header row over the cells, numbers kept as numbers."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    # Sheet names are capped at 31 characters and can't hold []:*?/\
    sheet.title = re.sub(r"[\[\]:*?/\\]", " ", title.strip() or "Sheet")[:31].strip() or "Sheet"
    book.properties.title = title
    sheet.append(columns)
    for row in rows:
        sheet.append([None if v == "" else v for v in row])
    fill = PatternFill("solid", start_color=HEADER_FILL, end_color=HEADER_FILL)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="000000")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    for i, name in enumerate(columns, start=1):
        longest = max([len(str(name))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r) and r[i - 1] is not None])
        sheet.column_dimensions[get_column_letter(i)].width = min(60, max(10, longest + 2))
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def to_pptx(title: str, slides: list[dict[str, str]]) -> bytes:
    """16:9 deck, one "Title and Content" slide each. Lines starting with "- " become bullets."""
    from pptx import Presentation
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches as PptInches

    deck = Presentation()
    deck.slide_width, deck.slide_height = PptInches(13.333), PptInches(7.5)
    deck.core_properties.title = title
    layout = deck.slide_layouts[1]
    for data in slides or [{"title": "", "body": ""}]:
        slide = deck.slides.add_slide(layout)
        heading, content = slide.placeholders[0], slide.placeholders[1]
        # The template is 4:3; stretch its boxes across the wider slide.
        for shape in (heading, content):
            shape.left, shape.width = PptInches(0.75), deck.slide_width - PptInches(1.5)
        heading.text = data.get("title", "")
        lines = [line for line in data.get("body", "").splitlines() if line.strip()]
        frame = content.text_frame
        if not lines:
            content.element.getparent().remove(content.element)
            continue
        for i, line in enumerate(lines):
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            bullet = line.lstrip().startswith(("- ", "* "))
            para.text = line.lstrip()[2:] if bullet else line.strip()
            if not bullet:  # plain lines: no bullet glyph
                pPr = para._p.get_or_add_pPr()
                pPr.set("indent", "0")
                pPr.set("marL", "0")
                pPr.insert(0, OxmlElement("a:buNone"))
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()
