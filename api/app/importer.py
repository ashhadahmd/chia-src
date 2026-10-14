"""Turn uploaded files and Google Workspace files into Chia documents.

Word (.docx), Markdown and text become docs; Excel (.xlsx) and CSV become sheets (one per
worksheet); PowerPoint (.pptx) becomes slides. Google Docs/Sheets/Slides are fetched as
.docx/.xlsx/.pptx and go through the same converters.

Chia's editor keeps formatting simple, so imports keep what it can show: headings, lists,
bold/italic/underline, font size, alignment, page size and text colour (snapped to Chia's
palette so it stays readable in dark mode). Images, charts and tables-in-docs are reduced or
skipped, and each import reports what it left out.
"""

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from html import escape
from pathlib import PurePath
from typing import Any
from urllib.parse import unquote

import httpx

from app.documents import DocumentError, markdown_to_html
from app.models import DocKind, PageSize

MAX_UPLOAD = 25 * 1024 * 1024
MAX_UNZIPPED = 300 * 1024 * 1024  # zip-bomb guard for .docx/.xlsx/.pptx
MAX_ROWS = 5000
MAX_COLS = 100

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

LEGACY = {
    ".doc": ".docx",
    ".xls": ".xlsx",
    ".ppt": ".pptx",
    ".odt": ".docx",
    ".ods": ".xlsx",
    ".odp": ".pptx",
    ".rtf": ".docx",
}
ACCEPTED = (".docx", ".xlsx", ".xlsm", ".csv", ".tsv", ".pptx", ".md", ".markdown", ".txt")


@dataclass
class Imported:
    kind: DocKind
    title: str
    content: dict[str, Any]


@dataclass
class Result:
    items: list[Imported] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


def convert(filename: str, data: bytes) -> Result:
    path = PurePath(filename or "Untitled")
    ext, title = path.suffix.lower(), path.stem.strip() or "Untitled"
    if ext in LEGACY:
        raise DocumentError(415, f"{ext} files can't be read directly. Save it as {LEGACY[ext]} and import that.")
    if not data:
        raise DocumentError(422, "That file is empty")
    try:
        match ext:
            case ".docx":
                return _docx(title, _office(data, ext))
            case ".xlsx" | ".xlsm":
                return _xlsx(title, _office(data, ext))
            case ".pptx":
                return _pptx(title, _office(data, ext))
            case ".csv" | ".tsv":
                return _csv(title, data, "\t" if ext == ".tsv" else None)
            case ".md" | ".markdown":
                return Result([Imported(DocKind.doc, title, {"html": markdown_to_html(_text(data)), "page": "letter"})])
            case ".txt":
                paragraphs = "".join(f"<p>{escape(line)}</p>" for line in _text(data).splitlines())
                return Result([Imported(DocKind.doc, title, {"html": paragraphs or "<p></p>", "page": "letter"})])
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError(422, f"Couldn't read “{path.name}”. It may be damaged or password-protected.") from exc
    raise DocumentError(415, f"Chia can't import {ext or 'that kind of'} files yet. Try Word, Excel, PowerPoint, CSV or Markdown.")


def _office(data: bytes, ext: str) -> bytes:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise DocumentError(422, f"This isn't a valid {ext} file. It may be damaged or password-protected.") from exc
    if sum(i.file_size for i in archive.infolist()) > MAX_UNZIPPED:
        raise DocumentError(413, "That file expands to more than Chia can import")
    return data


def _text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _plural(n: int, word: str) -> str:
    return f"{n} {word}{'' if n == 1 else 's'}"


# ───────────── Colours ─────────────

# Must match TEXT_COLORS in app/export.py and web/src/components/DocEditor.tsx (light hexes).
_PALETTE = {
    "stone": (0x6F, 0x6E, 0x66),
    "sage": (0x5A, 0x6B, 0x47),
    "eucalyptus": (0x3F, 0x6E, 0x66),
    "terracotta": (0xB0, 0x5A, 0x47),
    "ochre": (0x9A, 0x74, 0x24),
    "clay": (0x8C, 0x5A, 0x3C),
    "plum": (0x7A, 0x4F, 0x72),
    "slate": (0x4A, 0x5A, 0x6E),
}


def _palette_color(rgb: tuple[int, int, int]) -> str | None:
    """Snap a document colour to Chia's palette; near-black stays the default ink."""
    r, g, b = rgb
    if max(rgb) - min(rgb) < 24 and max(rgb) < 90:
        return None
    name = min(_PALETTE, key=lambda n: sum((a - c) ** 2 for a, c in zip(_PALETTE[n], (r, g, b), strict=True)))
    p = _PALETTE[name]
    return f"var(--tc-{name}, #{p[0]:02X}{p[1]:02X}{p[2]:02X})"


# ───────────── Word ─────────────

DEFAULT_PT = 11


@dataclass(frozen=True)
class _Style:
    bold: bool = False
    italic: bool = False
    underline: bool = False
    size: float | None = None
    color: str | None = None


def _inline(segments: list[tuple[str, _Style]]) -> str:
    merged: list[tuple[str, _Style]] = []
    for text, style in segments:
        if merged and merged[-1][1] == style:
            merged[-1] = (merged[-1][0] + text, style)
        elif text:
            merged.append((text, style))
    html = []
    for text, s in merged:
        out = escape(text).replace("\n", "<br>").replace("\t", "&emsp;")
        css = []
        if s.size and round(s.size) != DEFAULT_PT:
            css.append(f"font-size: {round(s.size)}pt")
        if s.color:
            css.append(f"color: {s.color}")
        if css:
            out = f'<span style="{"; ".join(css)}">{out}</span>'
        if s.underline:
            out = f"<u>{out}</u>"
        if s.italic:
            out = f"<em>{out}</em>"
        if s.bold:
            out = f"<strong>{out}</strong>"
        html.append(out)
    return "".join(html)


def _docx(title: str, data: bytes) -> Result:
    from docx import Document as load_docx
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.hyperlink import Hyperlink

    doc = load_docx(io.BytesIO(data))
    notes: list[str] = []
    images = tables = 0
    normal = doc.styles.default(WD_STYLE_TYPE.PARAGRAPH)
    base_pt = normal.font.size.pt if normal is not None and normal.font.size else DEFAULT_PT

    def numbering_format(p: Any) -> str | None:
        """'bullet', a number format, or None when the paragraph isn't in a list."""
        p_pr = p._p.pPr
        num_pr = p_pr.numPr if p_pr is not None else None
        style_name = (p.style.name if p.style is not None else "") or ""
        if num_pr is None or num_pr.numId is None:
            if style_name.startswith("List Bullet"):
                return "bullet"
            if style_name.startswith("List Number"):
                return "decimal"
            return None
        num_id, level = num_pr.numId.val, num_pr.ilvl.val if num_pr.ilvl is not None else 0
        if num_id == 0:
            return None
        try:
            numbering = doc.part.numbering_part.element
            abstract_id = numbering.num_having_numId(num_id).abstractNumId.val
            fmt = numbering.xpath(
                f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]/w:lvl[@w:ilvl="{level}"]/w:numFmt/@w:val'
            )
            return fmt[0] if fmt else "bullet"
        except Exception:
            return "decimal" if "Number" in style_name else "bullet"

    def run_style(run: Any, heading: bool) -> _Style:
        size = run.font.size.pt if run.font.size else None
        color = None
        try:
            if run.font.color is not None and run.font.color.type is not None and run.font.color.rgb is not None:
                rgb = run.font.color.rgb
                color = _palette_color((rgb[0], rgb[1], rgb[2]))
        except (AttributeError, ValueError, TypeError):
            color = None
        return _Style(
            bold=bool(run.bold) and not heading,
            italic=bool(run.italic),
            underline=bool(run.underline),
            size=None if heading else (size or base_pt),
            color=color,
        )

    def paragraph_html(p: Any, heading: bool) -> str:
        nonlocal images
        segments: list[tuple[str, _Style]] = []
        for item in p.iter_inner_content():
            runs = item.runs if isinstance(item, Hyperlink) else [item]
            for run in runs:
                images += len(run._r.findall(f".//{qn('w:drawing')}")) + len(run._r.findall(f".//{qn('w:pict')}"))
                segments.append((run.text, run_style(run, heading)))
        return _inline(segments)

    def alignment_attr(p: Any) -> str:
        align = p.alignment if p.alignment is not None else (p.style.paragraph_format.alignment if p.style is not None else None)
        value = {
            WD_ALIGN_PARAGRAPH.CENTER: "center",
            WD_ALIGN_PARAGRAPH.RIGHT: "right",
            WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
            WD_ALIGN_PARAGRAPH.DISTRIBUTE: "justify",
        }.get(align)
        return f' style="text-align: {value}"' if value else ""

    blocks: list[str] = []
    open_list: str | None = None

    def close_list() -> None:
        nonlocal open_list
        if open_list:
            blocks.append(f"</{open_list}>")
            open_list = None

    for block in doc.iter_inner_content():
        if isinstance(block, Table):
            tables += 1
            close_list()
            for row in block.rows:
                cells: list[str] = []
                for cell in row.cells:  # merged cells repeat; keep each once
                    text = cell.text.strip()
                    if text and (not cells or cells[-1] != text):
                        cells.append(text)
                if cells:
                    blocks.append(f"<p>{escape('  ·  '.join(cells))}</p>")
            continue

        style_name = (block.style.name if block.style is not None else "") or ""
        level = 0
        if style_name == "Title":
            level = 1
        elif m := re.fullmatch(r"Heading (\d)", style_name):
            level = min(int(m.group(1)), 3)
        fmt = None if level else numbering_format(block)
        body = paragraph_html(block, heading=bool(level))

        if fmt:
            tag = "ul" if fmt in ("bullet", "none") else "ol"
            if open_list != tag:
                close_list()
                blocks.append(f"<{tag}>")
                open_list = tag
            blocks.append(f"<li><p{alignment_attr(block)}>{body}</p></li>")
            continue
        close_list()
        tag = f"h{level}" if level else "p"
        blocks.append(f"<{tag}{alignment_attr(block)}>{body}</{tag}>")
    close_list()

    page = PageSize.letter
    if doc.sections and doc.sections[0].page_width is not None:
        width_in = doc.sections[0].page_width.inches
        if abs(width_in - 8.27) < 0.15 or abs(width_in - 11.69) < 0.15:
            page = PageSize.a4

    if images:
        notes.append(f"{_plural(images, 'image')} left out — Chia docs are text only")
    if tables:
        notes.append(f"{_plural(tables, 'table')} turned into plain lines")
    return Result([Imported(DocKind.doc, title, {"html": "".join(blocks) or "<p></p>", "page": page})], notes)


# ───────────── Sheets ─────────────


def _column_letter(i: int) -> str:
    name = ""
    i += 1
    while i:
        i, rem = divmod(i - 1, 26)
        name = chr(65 + rem) + name
    return name


def _cell(value: Any) -> str | int | float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int | float):
        return int(value) if isinstance(value, float) and value.is_integer() and abs(value) < 2**53 else value
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, date | time):
        return value.isoformat()
    text = str(value)
    return text if text.strip() else None


def _table(raw: list[list[Any]]) -> dict[str, Any] | None:
    """Trim empty edges and pick a header row. None when nothing is left."""
    rows = [[_cell(v) for v in r[:MAX_COLS]] for r in raw]
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    while rows and all(v is None for v in rows[0]):
        rows.pop(0)
    if not rows:
        return None
    width = max((max((i + 1 for i, v in enumerate(r) if v is not None), default=0) for r in rows), default=0)
    if width == 0:
        return None
    rows = [(r + [None] * width)[:width] for r in rows]
    first = rows[0]
    if all(v is None or isinstance(v, str) for v in first) and sum(v is not None for v in first) >= max(1, width // 2):
        columns = [str(v) if v is not None else _column_letter(i) for i, v in enumerate(first)]
        rows = rows[1:]
    else:
        columns = [_column_letter(i) for i in range(width)]
    return {"columns": columns, "rows": rows or [[None] * width]}


def _xlsx(title: str, data: bytes) -> Result:
    from openpyxl import load_workbook

    # data_only: imported sheets show each formula's last calculated value.
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    result = Result()
    truncated: list[str] = []
    sheets = book.worksheets
    for ws in sheets:
        raw: list[list[Any]] = []
        for row in ws.iter_rows(values_only=True, max_col=MAX_COLS):
            if len(raw) > MAX_ROWS:
                truncated.append(ws.title)
                break
            raw.append(list(row))
        table = _table(raw[:MAX_ROWS + 1])
        if table is None:
            continue
        name = title if len(sheets) == 1 else f"{title} · {ws.title}"
        result.items.append(Imported(DocKind.sheet, name, table))
    charts = len(book.chartsheets)
    book.close()
    if not result.items:
        raise DocumentError(422, "That workbook has no data to import")
    if charts:
        result.notes.append(f"{_plural(charts, 'chart sheet')} left out")
    if truncated:
        result.notes.append(f"Only the first {MAX_ROWS:,} rows of {', '.join(truncated)} came across")
    return result


def _csv(title: str, data: bytes, delimiter: str | None) -> Result:
    text = _text(data)
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    raw: list[list[Any]] = []
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        if len(raw) > MAX_ROWS:
            break
        raw.append(list(row))
    table = _table(raw[: MAX_ROWS + 1])
    if table is None:
        raise DocumentError(422, "That file has no data to import")
    notes = [f"Only the first {MAX_ROWS:,} rows came across"] if len(raw) > MAX_ROWS else []
    return Result([Imported(DocKind.sheet, title, table)], notes)


# ───────────── Slides ─────────────


def _pptx(title: str, data: bytes) -> Result:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

    deck = Presentation(io.BytesIO(data))
    title_types = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
    bullet_types = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.VERTICAL_BODY, PP_PLACEHOLDER.VERTICAL_OBJECT}
    skipped = 0

    def shapes(group: Any) -> list[Any]:
        found = []
        for shape in group:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                found.extend(shapes(shape.shapes))
            else:
                found.append(shape)
        return found

    def placeholder_type(shape: Any) -> Any:
        return shape.placeholder_format.type if shape.is_placeholder else None

    slides = []
    for slide in deck.slides:
        slide_title = ""
        lines: list[str] = []
        ordered = sorted(shapes(slide.shapes), key=lambda s: (s.top or 0, s.left or 0))
        for shape in ordered:
            kind = placeholder_type(shape)
            if shape.has_text_frame:
                text_frame = shape.text_frame
                if kind in title_types and not slide_title:
                    slide_title = " ".join(text_frame.text.split())
                    continue
                bulleted = kind in bullet_types
                for para in text_frame.paragraphs:
                    text = " ".join(para.text.replace("\v", " ").split())
                    if text:
                        lines.append(f"- {text}" if bulleted or para.level > 0 else text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [" ".join(c.text.split()) for c in row.cells]
                    if any(cells):
                        lines.append("  |  ".join(cells))
            elif shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.MEDIA) or getattr(
                shape, "has_chart", False
            ):
                skipped += 1
        slides.append({"title": slide_title, "body": "\n".join(lines)})

    notes = [f"{_plural(skipped, 'image')} and charts left out — Chia slides are text only"] if skipped else []
    return Result([Imported(DocKind.slides, title, {"slides": slides or [{"title": title, "body": ""}]})], notes)


# ───────────── Google links ─────────────

_GOOGLE_LINK = re.compile(
    r"^https?://(?:docs|drive)\.google\.com/(?:(document|spreadsheets|presentation)/(?:u/\d+/)?d/|file/(?:u/\d+/)?d/|open\?id=)"
    r"([A-Za-z0-9_-]{15,})"
)
# Google editor type → (export format, extension)
GOOGLE_EXPORTS = {"document": ("docx", ".docx"), "spreadsheets": ("xlsx", ".xlsx"), "presentation": ("pptx", ".pptx")}
_GOOGLE_HOSTS = (".google.com", ".googleusercontent.com")


def parse_google_link(url: str) -> tuple[str | None, str]:
    """Return (editor type or None for plain Drive files, file id)."""
    match = _GOOGLE_LINK.match(url.strip())
    if not match:
        raise DocumentError(422, "Paste a link to a Google Doc, Sheet or Slides file")
    return match.group(1), match.group(2)


def _filename_from(disposition: str) -> str | None:
    if m := re.search(r"filename\*=UTF-8''([^;]+)", disposition):
        return unquote(m.group(1))
    if m := re.search(r'filename="([^"]+)"', disposition):
        return m.group(1)
    return None


def fetch_public(editor: str, file_id: str) -> tuple[str, bytes]:
    """Download a link-shared Google file without any Google sign-in."""
    fmt, ext = GOOGLE_EXPORTS[editor]
    url = f"https://docs.google.com/{editor}/d/{file_id}/export?format={fmt}"
    private = DocumentError(
        403,
        "Google wouldn't share that file. Set its sharing to “Anyone with the link”, "
        "or connect Google and pick it from Drive.",
    )
    with httpx.Client(follow_redirects=True, max_redirects=5, timeout=30) as client:
        with client.stream("GET", url) as res:
            host = res.url.host or ""
            if not host.endswith(_GOOGLE_HOSTS):
                raise DocumentError(502, "Google sent Chia somewhere unexpected")
            if res.status_code in (401, 403, 404) or res.headers.get("content-type", "").startswith("text/html"):
                raise private
            if res.status_code >= 400:
                raise DocumentError(502, f"Google returned an error ({res.status_code})")
            chunks, size = [], 0
            for chunk in res.iter_bytes():
                size += len(chunk)
                if size > MAX_UPLOAD:
                    raise DocumentError(413, "That file is larger than 25 MB")
                chunks.append(chunk)
            name = _filename_from(res.headers.get("content-disposition", "")) or f"Untitled{ext}"
    return name, b"".join(chunks)
