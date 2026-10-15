"""Word (.docx) rendering of editor HTML; see app.export. Imported only when a doc is exported,
so python-docx stays out of the API's cold start."""

import io
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from typing import Any

from docx import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Mm, Pt, RGBColor

from app.assets import ASSET_SRC, Resolver
from app.export import HEADER_FILL, _no_resolver
from app.models import PageSize

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_MARKS = {"strong": "bold", "b": "bold", "em": "italic", "i": "italic", "u": "underline", "s": "strike", "strike": "strike", "code": "code"}
_VOID = {"br", "hr", "img", "col", "input", "meta", "link", "source", "wbr"}
_DPI = 96


@dataclass
class _Node:
    tag: str
    attrs: dict[str, str] = field(default_factory=dict)
    children: list[_Node | str] = field(default_factory=list)

    @property
    def style(self) -> dict[str, str]:
        raw = self.attrs.get("style") or ""
        return {k.strip().lower(): v.strip() for k, v in (d.split(":", 1) for d in raw.split(";") if ":" in d)}

    def elements(self, *tags: str) -> list[_Node]:
        return [c for c in self.children if isinstance(c, _Node) and (not tags or c.tag in tags)]

    def text(self) -> str:
        return "".join(c if isinstance(c, str) else c.text() for c in self.children)


class _TreeBuilder(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.root = _Node("root")
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs) -> None:
        node = _Node(tag, {k: v or "" for k, v in attrs})
        self.stack[-1].children.append(node)
        if tag not in _VOID:
            self.stack.append(node)

    def handle_startendtag(self, tag, attrs) -> None:
        self.stack[-1].children.append(_Node(tag, {k: v or "" for k, v in attrs}))

    def handle_endtag(self, tag) -> None:
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                return

    def handle_data(self, data) -> None:
        self.stack[-1].children.append(data)


def parse_tree(html: str) -> _Node:
    builder = _TreeBuilder()
    builder.feed(html)
    builder.close()
    return builder.root


def _points(value: str) -> float | None:
    m = re.match(r"^([\d.]+)\s*(pt|px)?$", value)
    if not m:
        return None
    n = float(m[1])
    return n * 0.75 if m[2] == "px" else n


def _int(value: str | None, default: int = 1) -> int:
    try:
        return max(1, int(value or default))
    except ValueError:
        return default


class _Box:
    """Where blocks go: the document body or a table cell (which starts with an empty paragraph)."""

    def __init__(self, target: Any, width: int, fresh_cell: bool = False) -> None:
        self.target = target
        self.width = width  # EMU available for content
        self.spare = target.paragraphs[0] if fresh_cell else None

    def paragraph(self, style: str | None = None):
        if self.spare is not None:
            para, self.spare = self.spare, None
            if style:
                para.style = style
            return para
        return self.target.add_paragraph(style=style)


class _DocxWriter:
    def __init__(self, doc: Any, resolve: Resolver) -> None:
        self.doc = doc
        self.resolve = resolve

    # Blocks

    def blocks(self, box: _Box, nodes: list[_Node | str], lists: tuple[str, ...] = (), quote: bool = False) -> None:
        loose: list[_Node | str] = []  # inline content sitting directly in a block container

        def flush() -> None:
            if any(isinstance(n, _Node) or n.strip() for n in loose):
                para = box.paragraph(self._para_style(None, lists, quote))
                for n in loose:
                    self.inline(para, n, [])
            loose.clear()

        for node in nodes:
            if isinstance(node, str) or node.tag in _MARKS or node.tag in ("span", "a", "br"):
                loose.append(node)
                continue
            flush()
            tag = node.tag
            if tag in ("p", "h1", "h2", "h3"):
                para = box.paragraph(self._para_style(tag, lists, quote))
                if align := _ALIGN.get(node.style.get("text-align", "")):
                    para.alignment = align
                for child in node.children:
                    self.inline(para, child, [])
            elif tag in ("ul", "ol"):
                for li in node.elements("li"):
                    self.blocks(box, li.children, (*lists, tag), quote)
            elif tag == "blockquote":
                self.blocks(box, node.children, lists, True)
            elif tag == "hr":
                para = box.paragraph()
                para.add_run("─" * 30)
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif tag == "img":
                self.image(box, node, node)
            elif tag == "figure":
                img = next((c for c in node.elements() if c.tag == "img"), None)
                if img is not None:
                    self.image(box, img, node)
                elif code := node.attrs.get("data-code"):
                    self.code(box, code)
                else:
                    self.blocks(box, node.children, lists, quote)
            elif tag == "pre":
                self.code(box, node.text())
            elif tag == "table":
                self.table(box, node)
            else:
                self.blocks(box, node.children, lists, quote)
        flush()

    @staticmethod
    def _para_style(tag: str | None, lists: tuple[str, ...], quote: bool) -> str | None:
        if tag in ("h1", "h2", "h3"):
            return f"Heading {tag[1]}"
        if lists:
            base = "List Bullet" if lists[-1] == "ul" else "List Number"
            depth = min(len(lists), 3)
            return base if depth == 1 else f"{base} {depth}"
        return "Quote" if quote else None

    def code(self, box: _Box, text: str) -> None:
        para = box.paragraph()
        for i, line in enumerate(text.split("\n")):
            if i:
                para.add_run().add_break()
            run = para.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9)

    def image(self, box: _Box, img: _Node, holder: _Node) -> None:
        m = ASSET_SRC.search(img.attrs.get("src", ""))
        asset = self.resolve(m.group(1)) if m else None
        if asset is None:
            if alt := img.attrs.get("alt"):
                box.paragraph().add_run(f"[{alt}]").italic = True
            return
        info = DocxImage.from_blob(asset.data)
        natural = Emu(int(info.px_width / _DPI * 914400)) if info.px_width else box.width
        try:
            percent = float(holder.attrs.get("data-width") or img.attrs.get("data-width") or 0)
        except ValueError:
            percent = 0
        width = int(box.width * min(percent, 100) / 100) if percent > 0 else min(natural, box.width)
        para = box.paragraph()
        para.alignment = _ALIGN.get(holder.attrs.get("data-align") or img.attrs.get("data-align") or "center")
        para.add_run().add_picture(io.BytesIO(asset.data), width=Emu(width))

    def table(self, box: _Box, node: _Node) -> None:
        rows: list[_Node] = []
        for part in node.elements():
            if part.tag == "tr":
                rows.append(part)
            elif part.tag in ("thead", "tbody", "tfoot"):
                rows.extend(part.elements("tr"))
        if not rows:
            return

        # Lay cells out on a grid, honouring colspan/rowspan.
        taken: set[tuple[int, int]] = set()
        placed: list[tuple[int, int, int, int, _Node]] = []
        width_px: dict[int, float] = {}
        for r, tr in enumerate(rows):
            c = 0
            for cell in tr.elements("td", "th"):
                while (r, c) in taken:
                    c += 1
                cs, rs = _int(cell.attrs.get("colspan")), _int(cell.attrs.get("rowspan"))
                rs = min(rs, len(rows) - r)
                for dr in range(rs):
                    for dc in range(cs):
                        taken.add((r + dr, c + dc))
                widths = [w for w in (cell.attrs.get("colwidth") or cell.attrs.get("data-colwidth") or "").split(",") if w.strip()]
                for i, w in enumerate(widths[:cs]):
                    try:
                        if float(w) > 0:
                            width_px.setdefault(c + i, float(w))
                    except ValueError:
                        pass
                placed.append((r, c, rs, cs, cell))
                c += cs
        ncols = max(c for _, c in taken) + 1
        table = box.target.add_table(rows=len(rows), cols=ncols)
        table.style = self.doc.styles["Table Grid"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Column widths: as in the editor, sized columns keep their pixel width and the rest share
        # what is left of the row; scaled down if that overflows the page.
        page_px = box.width / 914400 * _DPI
        known = sum(width_px.values())
        unknown = ncols - len(width_px)
        default_px = max(48.0, (page_px - known) / unknown) if unknown else 0
        px_total = known + unknown * default_px
        scale = min(box.width / max(px_total, 1), 914400 / _DPI)
        col_emu = [int((width_px.get(i, default_px)) * scale) for i in range(ncols)]
        for i, col in enumerate(table.columns):
            col.width = Emu(col_emu[i])
            for cell in col.cells:
                cell.width = Emu(col_emu[i])

        header_rows = {r for r, tr in enumerate(rows) if tr.elements("th") and len(tr.elements("th")) == len(tr.elements("td", "th"))}
        for r in header_rows:
            if r == 0 or r - 1 in header_rows:
                tr_pr = table.rows[r]._tr.get_or_add_trPr()
                tr_pr.append(OxmlElement("w:tblHeader"))

        for r, c, rs, cs, node_cell in placed:
            cell = table.cell(r, c)
            if rs > 1 or cs > 1:
                cell = cell.merge(table.cell(r + rs - 1, c + cs - 1))
            inner = _Box(cell, sum(col_emu[c : c + cs]) - int(Inches(0.15)), fresh_cell=True)
            self.blocks(inner, node_cell.children)
            if node_cell.tag == "th":
                shading = OxmlElement("w:shd")
                shading.set(qn("w:val"), "clear")
                shading.set(qn("w:color"), "auto")
                shading.set(qn("w:fill"), HEADER_FILL)
                cell._tc.get_or_add_tcPr().append(shading)
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(0, 0, 0)
        # Word glues consecutive tables together; a paragraph keeps them apart.
        box.target.add_paragraph()

    # Inline

    def inline(self, para: Any, node: _Node | str, marks: list[dict[str, Any]]) -> None:
        if isinstance(node, str):
            if not node:
                return
            run = para.add_run(node)
            for mark in marks:
                if mark.get("bold"):
                    run.bold = True
                if mark.get("italic"):
                    run.italic = True
                if mark.get("underline"):
                    run.underline = True
                if mark.get("strike"):
                    run.font.strike = True
                if mark.get("code"):
                    run.font.name = "Consolas"
                if mark.get("size"):
                    run.font.size = Pt(mark["size"])
                if mark.get("color"):
                    run.font.color.rgb = RGBColor.from_string(mark["color"].upper())
            return
        if node.tag == "br":
            para.add_run().add_break()
            return
        mark: dict[str, Any] = {}
        if node.tag in _MARKS:
            mark[_MARKS[node.tag]] = True
        elif node.tag == "a":
            mark["underline"] = True
        elif node.tag == "span":
            style = node.style
            if size := _points(style.get("font-size", "")):
                mark["size"] = size
            if hex_match := re.search(r"#([0-9a-fA-F]{6})", style.get("color", "")):
                mark["color"] = hex_match[1]
        for child in node.children:
            self.inline(para, child, [*marks, mark])


def to_docx(title: str, html: str, page: PageSize, resolve: Resolver = _no_resolver) -> bytes:
    doc = DocxDocument()
    doc.core_properties.title = title
    section = doc.sections[0]
    if page is PageSize.letter:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    else:
        section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(1))

    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.5

    width = section.page_width - section.left_margin - section.right_margin
    _DocxWriter(doc, resolve).blocks(_Box(doc, width), parse_tree(html).children)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


